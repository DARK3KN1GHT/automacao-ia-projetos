import os
import io
import pandas as pd
import streamlit as st
from openai import OpenAI
from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Carrega as variáveis de ambiente
load_dotenv()
api_key = os.environ.get("GROQ_API_KEY")

client = OpenAI(
    base_url="https://api.groq.com/openai/v1",
    api_key=api_key,
)

st.set_page_config(page_title="Automação de Dados com IA", page_icon="📊", layout="wide")

st.title("📊 Painel de Automação de Dados e Análise Inteligente")
st.write("Carregue o seu ficheiro CSV para processar dados, personalizar relatórios com IA e exportar em Markdown e Excel Profissional.")

# --- BARRA LATERAL (HISTÓRICO E CONFIGURAÇÕES) ---
st.sidebar.header("📁 Histórico de Relatórios")
if os.path.exists("relatorio_executivo.md"):
    if st.sidebar.button("Ver Último Relatório Guardado"):
        with open("relatorio_executivo.md", "r", encoding="utf-8") as f:
            conteudo_historico = f.read()
        st.sidebar.markdown(conteudo_historico)
else:
    st.sidebar.info("Nenhum relatório anterior encontrado.")

st.sidebar.divider()
st.sidebar.header("⚙️ Configurações da IA")
tom_relatorio = st.sidebar.selectbox(
    "Escolha o tom do relatório:",
    ["Executivo (Padrão)", "Comercial / Foco em Vendas", "Técnico / Foco em Custos"]
)

# Se o utilizador mudar o tom, limpamos o relatório anterior da sessão para evitar conflitos
if "tom_anterior" not in st.session_state:
    st.session_state["tom_anterior"] = tom_relatorio

if st.session_state["tom_anterior"] != tom_relatorio:
    st.session_state["tom_anterior"] = tom_relatorio
    if "relatorio_atual" in st.session_state:
        del st.session_state["relatorio_atual"] # Reseta a pré-visualização ao mudar de tom

# Mapeamento dos tons para os prompts do sistema
prompts_sistema = {
    "Executivo (Padrão)": "Você é um analista de dados sénior especialista em relatórios executivos.",
    "Comercial / Foco em Vendas": "Você é um diretor comercial focado em estratégias de vendas, expansão de mercado e aumento de receita.",
    "Técnico / Foco em Custos": "Você é um auditor financeiro e de operações focado em otimização de stock, margens e redução de custos."
}

# --- CORPO DA APLICAÇÃO ---
ficheiro_carregado = st.file_uploader("Escolha um ficheiro CSV", type=["csv"])

if ficheiro_carregado is not None:
    try:
        # Lê o CSV enviado pelo utilizador
        df = pd.read_csv(ficheiro_carregado)
        
        # Validação robusta de colunas obrigatórias
        colunas_obrigatorias = ["Produto", "Categoria", "Preco_Unitario", "Quantidade_Vendida"]
        colunas_em_falta = [col for col in colunas_obrigatorias if col not in df.columns]
        
        if colunas_em_falta:
            st.error(f"[-] O ficheiro enviado não tem as colunas obrigatórias: {colunas_em_falta}")
        else:
            st.subheader("📋 Dados Originais Validados")
            st.dataframe(df)
            
            if st.button("Executar Análise com IA e Gerar Pré-visualização"):
                with st.spinner(f"A processar dados com o tom '{tom_relatorio}'..."):
                    # Processamento com Pandas
                    df["Faturamento_Total"] = df["Preco_Unitario"] * df["Quantidade_Vendida"]
                    dados_em_texto = df.to_csv(index=False)
                    
                    # Chamada à API da Groq com o tom selecionado
                    system_prompt = prompts_sistema[tom_relatorio]
                    response = client.chat.completions.create(
                        model="openai/gpt-oss-safeguard-20b",
                        messages=[
                            {"role": "system", "content": system_prompt},
                            {"role": "user", "content": f"Elabore um relatório detalhado com base nestes dados:\n\n{dados_em_texto}"}
                        ],
                    )
                    
                    relatorio_ia = response.choices[0].message.content
                    
                    # Guarda na sessão do Streamlit
                    st.session_state["relatorio_atual"] = relatorio_ia
                    st.session_state["df_processado"] = df
                    
                    # Guarda também no disco para o histórico lateral
                    with open("relatorio_executivo.md", "w", encoding="utf-8") as f:
                        f.write(f"# Relatório Executivo ({tom_relatorio})\n\n")
                        f.write(relatorio_ia)
                
                st.success("Análise e pré-visualização geradas com sucesso!")
                st.balloons()

            # Pré-visualização antes do download
            if "relatorio_atual" in st.session_state:
                st.divider()
                st.header("🔍 Pré-visualização Antes do Download")
                
                aba_texto, aba_graficos = st.tabs(["📄 Pré-visualização do Relatório", "📈 Pré-visualização de Gráficos"])
                
                with aba_texto:
                    st.info(f"Conteúdo atual gerado com o tom: **{tom_relatorio}**")
                    st.markdown(st.session_state["relatorio_atual"])
                
                with aba_graficos:
                    st.info("Validação visual do faturamento total por produto:")
                    df_graf = st.session_state["df_processado"]
                    if "Produto" in df_graf.columns and "Faturamento_Total" in df_graf.columns:
                        st.bar_chart(df_graf.set_index("Produto")["Faturamento_Total"])
                
                st.divider()
                st.subheader("📥 Opções de Download Profissional")
                
                col1, col2 = st.columns(2)
                
                with col1:
                    st.download_button(
                        label="📄 Descarregar Relatório (.md)",
                        data=st.session_state["relatorio_atual"],
                        file_name="relatorio_executivo.md",
                        mime="text/markdown"
                    )
                
                with col2:
                    # Geração do Excel Profissional com openpyxl
                    df_excel = st.session_state["df_processado"]
                    output = io.BytesIO()
                    
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Relatório de Vendas"
                    ws.views.sheetView[0].showGridLines = True
                    
                    headers = list(df_excel.columns)
                    ws.append(headers)
                    
                    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                    data_font = Font(name="Calibri", size=11)
                    total_font = Font(name="Calibri", size=11, bold=True)
                    
                    thin_border = Border(
                        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
                    )
                    total_border = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))
                    
                    for col_num in range(1, len(headers) + 1):
                        cell = ws.cell(row=1, column=col_num)
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    
                    for row_idx, row in enumerate(df_excel.values, start=2):
                        ws.append(list(row))
                        for col_idx in range(1, len(row) + 1):
                            cell = ws.cell(row=row_idx, column=col_idx)
                            cell.font = data_font
                            cell.border = thin_border
                            if headers[col_idx - 1] in ["Preco_Unitario", "Faturamento_Total"]:
                                cell.number_format = '"R$ "* #,##0.00'
                                cell.alignment = Alignment(horizontal="right")
                            elif headers[col_idx - 1] == "Quantidade_Vendida":
                                cell.number_format = '#,##0'
                                cell.alignment = Alignment(horizontal="center")
                            else:
                                cell.alignment = Alignment(horizontal="left")
                    
                    last_row = len(df_excel) + 1
                    total_row_idx = last_row + 1
                    ws.cell(row=total_row_idx, column=1, value="TOTAL")
                    ws.cell(row=total_row_idx, column=4, value=f"=SUM(D2:D{last_row})")
                    ws.cell(row=total_row_idx, column=5, value=f"=SUM(E2:E{last_row})")
                    
                    for col_idx in range(1, len(headers) + 1):
                        cell = ws.cell(row=total_row_idx, column=col_idx)
                        cell.font = total_font
                        cell.border = total_border
                        if headers[col_idx - 1] in ["Preco_Unitario", "Faturamento_Total"]:
                            cell.number_format = '"R$ "* #,##0.00'
                            cell.alignment = Alignment(horizontal="right")
                        elif headers[col_idx - 1] == "Quantidade_Vendida":
                            cell.number_format = '#,##0'
                            cell.alignment = Alignment(horizontal="center")
                    
                    for col in ws.columns:
                        max_len = max(len(str(cell.value or '')) for cell in col)
                        col_letter = get_column_letter(col[0].column)
                        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
                    
                    wb.save(output)
                    excel_data = output.getvalue()
                    
                    st.download_button(
                        label="📊 Descarregar Relatório Excel Profissional (.xlsx)",
                        data=excel_data,
                        file_name="relatorio_vendas_profissional.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                    
    except Exception as e:
        st.error(f"Ocorreu um erro ao processar o ficheiro: {e}")
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def gerar_excel_executivo(df_excel, total_faturamento, total_quantidade, ticket_medio):
    """
    Gera um workbook corporativo do openpyxl com múltiplas abas limpas e alinhadas:
    - Aba 1: Resumo & KPIs Consolidados
    - Aba 2: Dados Detalhados com fórmulas de soma automáticas
    """
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    
    # Estilos corporativos de alto padrão
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    
    # --- ABA 1: RESUMO DE KPIS ---
    ws_kpi = wb.active
    ws_kpi.title = "Resumo & KPIs"
    ws_kpi.views.sheetView[0].showGridLines = True
    
    ws_kpi.cell(row=1, column=1, value="PAINEL EXECUTIVO DE INDICADORES").font = title_font
    ws_kpi.append([]) 
    ws_kpi.append(["Indicador", "Valor Consolidado"])
    
    for col_num in range(1, 3):
        cell = ws_kpi.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    kpi_data = [
        ["Faturamento Total", total_faturamento],
        ["Volume Total Vendido", total_quantidade],
        ["Ticket Médio Global", ticket_medio]
    ]
    
    for r_idx, row_val in enumerate(kpi_data, start=4):
        ws_kpi.append(row_val)
        c1 = ws_kpi.cell(row=r_idx, column=1)
        c2 = ws_kpi.cell(row=r_idx, column=2)
        c1.font = data_font
        c1.border = thin_border
        c2.font = data_font
        c2.border = thin_border
        if "Faturamento" in row_val[0] or "Ticket" in row_val[0]:
            c2.number_format = '"R$ "* #,##0.00'
            c2.alignment = Alignment(horizontal="right")
        else:
            c2.number_format = '#,##0'
            c2.alignment = Alignment(horizontal="center")

    for col in ws_kpi.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_kpi.column_dimensions[col_letter].width = max(max_len + 6, 22)

    # --- ABA 2: DADOS DETALHADOS ---
    ws_data = wb.create_sheet(title="Dados Detalhados")
    ws_data.views.sheetView[0].showGridLines = True
    
    headers = list(df_excel.columns)
    ws_data.append(headers)
    
    for col_num in range(1, len(headers) + 1):
        cell = ws_data.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for row_idx, row in enumerate(df_excel.values, start=2):
        ws_data.append(list(row))
        for col_idx in range(1, len(row) + 1):
            cell = ws_data.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if headers[col_idx - 1] in ["Preco_Unitario", "Faturamento_Total"]:
                cell.number_format = '"R$ "* #,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif headers[col_idx - 1] == "Quantidade_Vendida":
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")
    
    last_row = len(df_excel) + 1
    total_row_idx = last_row + 1
    total_font = Font(name="Calibri", size=11, bold=True)
    total_border = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))
    
    ws_data.cell(row=total_row_idx, column=1, value="TOTAL")
    ws_data.cell(row=total_row_idx, column=4, value=f"=SUM(D2:D{last_row})")
    ws_data.cell(row=total_row_idx, column=5, value=f"=SUM(E2:E{last_row})")
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws_data.cell(row=total_row_idx, column=col_idx)
        cell.font = total_font
        cell.border = total_border
        if headers[col_idx - 1] in ["Preco_Unitario", "Faturamento_Total"]:
            cell.number_format = '"R$ "* #,##0.00'
            cell.alignment = Alignment(horizontal="right")
        elif headers[col_idx - 1] == "Quantidade_Vendida":
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="center")
    
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = max(max_len + 4, 15)
    
    wb.save(output)
    return output.getvalue()